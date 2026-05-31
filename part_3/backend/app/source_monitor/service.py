"""MYH source monitoring, download, transformation, and refresh services."""

from __future__ import annotations

import csv
import hashlib
import os
import re
import shutil
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Callable, Iterable
from urllib.parse import unquote, urljoin, urlparse
from urllib.request import Request, urlopen

import psycopg

from backend.app.database import get_database_url, open_connection
from backend.app.source_monitor.models import SourceFileCandidate
from backend.app.source_monitor import repositories as repo
from backend.scripts.load_curated_data import (
    APPLICATION_INSERT_SQL,
    DEFAULT_INDEXES_PATH,
    DEFAULT_RESET_SCHEMA_PATH,
    DEFAULT_SCHEMA_PATH,
    build_application_row,
    clean_text,
    insert_lookup_rows,
    read_rows,
    resolve_csv_path,
    to_int,
)

DEFAULT_SOURCE_PAGE_URL = "https://www.myh.se/yrkeshogskolan/resultat-ansokningsomgangar/resultat-for-program"
BACKEND_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_DOWNLOAD_DIR = BACKEND_ROOT / "runtime" / "source_files"
DEFAULT_PROCESSED_DIR = BACKEND_ROOT / "runtime" / "processed"
SUPPORTED_EXTENSIONS = {".xlsx", ".xlsm", ".xls", ".csv"}
DOWNLOAD_LINK_PATTERN = re.compile(r"<a\b[^>]*?href=[\"']([^\"']+)[\"'][^>]*>(.*?)</a>", re.IGNORECASE | re.DOTALL)
TAG_PATTERN = re.compile(r"<[^>]+>")
YEAR_PATTERN = re.compile(r"(?<!\d)(20\d{2})(?!\d)")

REQUIRED_RAW_FIELDS = {
    "diarienummer",
    "utbildningsnamn",
    "utbildningsomrade",
    "beslut",
    "lan",
    "kommun",
    "yh_poang",
    "studieform",
    "utbildningsanordnare",
    "huvudmannatyp",
}

HEADER_ALIASES: dict[str, tuple[str, ...]] = {
    "diarienummer": ("diarienummer", "dnr", "ärendenummer", "arendenummer", "ansökningsnummer", "ansokningsnummer"),
    "utbildningsnamn": ("utbildningsnamn", "utbildning", "namn", "utbildningens namn"),
    "utbildningsomrade": ("utbildningsområde", "utbildningsomrade", "område", "omrade", "utbildningsområde namn"),
    "beslut": ("beslut", "beslutsstatus", "beslutstyp"),
    "lan": ("län", "lan", "region"),
    "kommun": ("kommun", "kommuner", "ort"),
    "flera_kommuner": ("flera kommuner", "flera_kommuner"),
    "antal_kommuner": ("antal kommuner", "antal_kommuner"),
    "yh_poang": ("yh-poäng", "yh poäng", "yh_poang", "yh-poang", "poäng", "poang"),
    "studieform": ("studieform", "form"),
    "studietakt_procent": ("studietakt", "studietakt procent", "studietakt_procent", "studietakt %"),
    "examenstyp": ("examenstyp",),
    "utbildningsanordnare": ("utbildningsanordnare", "anordnare", "leverantör", "leverantor"),
    "huvudmannatyp": ("huvudmannatyp", "huvudman", "typ av huvudman"),
    "sokta_utbildningsomgangar": ("sökta utbildningsomgångar", "sokta utbildningsomgangar", "sökta omgångar"),
    "beviljade_utbildningsomgangar": ("beviljade utbildningsomgångar", "beviljade utbildningsomgangar", "beviljade omgångar"),
    "sun5_inriktning": ("sun5 inriktning", "sun5_inriktning"),
    "sun5_inriktning_namn": ("sun5 inriktning namn", "sun5_inriktning_namn"),
    "seqf_niva": ("seqf nivå", "seqf niva", "seqf_niva"),
    "smalt_yrkesomrade": ("smalt yrkesområde", "smalt yrkesomrade", "smalt_yrkesomrade"),
    "sokta_platser_totalt": ("sökta platser totalt", "sokta platser totalt", "sokta_platser_totalt"),
    "beviljade_platser_totalt": ("beviljade platser totalt", "beviljade_platser_totalt"),
    "sokta_platser_per_utbildningsomgang": (
        "sökta platser per utbildningsomgång",
        "sokta platser per utbildningsomgang",
        "sokta_platser_per_utbildningsomgang",
    ),
}


@dataclass(frozen=True)
class SourceMonitorConfig:
    """Environment-backed source-monitor configuration."""

    source_url: str = DEFAULT_SOURCE_PAGE_URL
    monitor_enabled: bool = False
    interval_minutes: int = 360
    run_on_startup: bool = False
    auto_import_enabled: bool = False
    download_dir: Path = DEFAULT_DOWNLOAD_DIR
    processed_dir: Path = DEFAULT_PROCESSED_DIR
    timeout_seconds: int = 20


@dataclass(frozen=True)
class PageFetchResponse:
    """Small HTTP response shape for source-page fetchers."""

    status_code: int
    body: str
    content_type: str | None = None


def _env_bool(name: str, default: bool = False) -> bool:
    value = os.getenv(name)
    if value is None:
        return default
    return value.strip().lower() in {"1", "true", "yes", "on"}


def _env_int(name: str, default: int) -> int:
    value = os.getenv(name)
    if not value:
        return default
    try:
        parsed = int(value)
        return parsed if parsed > 0 else default
    except ValueError:
        return default


def build_config_from_env() -> SourceMonitorConfig:
    """Build source-monitor configuration from environment variables."""
    return SourceMonitorConfig(
        source_url=os.getenv("PART3_MYH_SOURCE_PAGE_URL", DEFAULT_SOURCE_PAGE_URL),
        monitor_enabled=_env_bool("PART3_SOURCE_MONITOR_ENABLED", False),
        interval_minutes=_env_int("PART3_SOURCE_MONITOR_INTERVAL_MINUTES", 360),
        run_on_startup=_env_bool("PART3_SOURCE_MONITOR_RUN_ON_STARTUP", False),
        auto_import_enabled=_env_bool("PART3_REFRESH_AUTO_IMPORT", False),
        download_dir=Path(os.getenv("PART3_SOURCE_DOWNLOAD_DIR", str(DEFAULT_DOWNLOAD_DIR))),
        processed_dir=Path(os.getenv("PART3_PROCESSED_OUTPUT_DIR", str(DEFAULT_PROCESSED_DIR))),
        timeout_seconds=_env_int("PART3_SOURCE_MONITOR_TIMEOUT_SECONDS", 20),
    )


def fetch_source_page(source_url: str, timeout_seconds: int = 20) -> PageFetchResponse:
    """Fetch the MYH result page using the standard library."""
    request = Request(
        source_url,
        headers={"User-Agent": "myh-applications-api-source-monitor/0.3.20"},
        method="GET",
    )
    with urlopen(request, timeout=timeout_seconds) as response:
        body = response.read().decode("utf-8", errors="replace")
        return PageFetchResponse(response.getcode(), body, response.headers.get("Content-Type"))


def download_url_bytes(file_url: str, timeout_seconds: int = 30) -> bytes:
    """Download source-file bytes with a safe user agent."""
    request = Request(file_url, headers={"User-Agent": "myh-applications-api-refresh/0.3.20"}, method="GET")
    with urlopen(request, timeout=timeout_seconds) as response:
        return response.read()


def strip_html(value: str) -> str:
    """Remove tags and collapse whitespace from a small HTML label."""
    text = TAG_PATTERN.sub(" ", value)
    return " ".join(text.split())


def _file_name_from_url(file_url: str) -> str:
    parsed = urlparse(file_url)
    return unquote(Path(parsed.path).name)


def _file_type(file_name_or_url: str) -> str:
    suffix = Path(urlparse(file_name_or_url).path).suffix.lower()
    return suffix.lstrip(".")


def extract_source_year(*values: str | None) -> int | None:
    """Extract the latest visible 20xx source year from labels, names, or URLs."""
    years: set[int] = set()
    for value in values:
        if not value:
            continue
        years.update(int(match.group(1)) for match in YEAR_PATTERN.finditer(value))
    return max(years) if years else None


def parse_source_page(html: str, source_url: str) -> list[SourceFileCandidate]:
    """Parse official downloadable MYH source-file candidates from HTML.

    Unsupported and unrelated file types, such as PDFs, are ignored. The parser
    intentionally supports Excel first because the official MYH result page is
    workbook based, while CSV remains useful for local fixtures and internal
    processed outputs.
    """
    candidates: list[SourceFileCandidate] = []
    seen_urls: set[str] = set()
    for match in DOWNLOAD_LINK_PATTERN.finditer(html):
        href = match.group(1).strip()
        absolute_url = urljoin(source_url, href)
        file_name = _file_name_from_url(absolute_url)
        suffix = Path(urlparse(absolute_url).path).suffix.lower()
        if suffix not in SUPPORTED_EXTENSIONS:
            continue
        if absolute_url in seen_urls:
            continue
        seen_urls.add(absolute_url)
        label = strip_html(match.group(2)) or file_name
        candidates.append(
            SourceFileCandidate(
                file_name=file_name,
                file_url=absolute_url,
                file_type=_file_type(file_name),
                source_year=extract_source_year(file_name, absolute_url, label),
                label=label,
            )
        )
    return candidates


def _metadata_changed(existing: dict[str, Any], candidate: SourceFileCandidate) -> bool:
    """Return true when visible source metadata changed since last check."""
    return any(
        [
            existing.get("file_name") != candidate.file_name,
            existing.get("file_type") != candidate.file_type,
            existing.get("source_year") != candidate.source_year,
        ]
    )


def run_source_check(
    conn: psycopg.Connection[dict[str, Any]],
    *,
    actor_user_id: str | None = None,
    config: SourceMonitorConfig | None = None,
    fetcher: Callable[[str, int], PageFetchResponse] = fetch_source_page,
) -> dict[str, Any]:
    """Fetch the source page, store detected files, and notify admins."""
    config = config or build_config_from_env()
    check_run = repo.create_check_run(conn, source_url=config.source_url)
    check_run_id = str(check_run["id"])
    # Persist the check-run shell before fetching/parsing. If a later database
    # write fails, PostgreSQL aborts the current transaction; committing here
    # lets the error path roll back the failed write and still record the
    # failed check run plus an admin notification.
    conn.commit()
    files: list[dict[str, Any]] = []
    new_count = changed_count = known_count = 0

    try:
        response = fetcher(config.source_url, config.timeout_seconds)
        candidates = parse_source_page(response.body, config.source_url)
        for candidate in candidates:
            existing = repo.find_source_file_by_url(conn, candidate.file_url)
            if existing is None:
                stored = repo.create_source_file(
                    conn,
                    file_name=candidate.file_name,
                    file_url=candidate.file_url,
                    file_type=candidate.file_type,
                    source_year=candidate.source_year,
                    check_run_id=check_run_id,
                )
                new_count += 1
                repo.create_notification(
                    conn,
                    notification_type="source_file_detected",
                    severity="info",
                    title="New official MYH source file detected",
                    message=f"{candidate.file_name} was discovered on the MYH result page.",
                    source_file_id=stored["id"],
                    actor_user_id=actor_user_id,
                )
                stored["detection_status"] = "new"
            else:
                changed = _metadata_changed(existing, candidate)
                status = "changed" if changed else "known"
                stored = repo.update_seen_source_file(
                    conn,
                    source_file_id=existing["id"],
                    file_name=candidate.file_name,
                    file_type=candidate.file_type,
                    source_year=candidate.source_year,
                    status=status,
                    check_run_id=check_run_id,
                )
                if changed:
                    changed_count += 1
                    repo.create_notification(
                        conn,
                        notification_type="source_file_changed",
                        severity="warning",
                        title="Official MYH source file changed",
                        message=f"{candidate.file_name} changed visible metadata and should be reviewed before import.",
                        source_file_id=stored["id"],
                        actor_user_id=actor_user_id,
                    )
                    stored["detection_status"] = "changed"
                else:
                    known_count += 1
                    stored["detection_status"] = "known"
            files.append(stored)

        status = "success"
        message = f"Source check completed: {new_count} new, {changed_count} changed, {known_count} known file(s)."
        finished = repo.finish_check_run(
            conn,
            check_run_id=check_run_id,
            status=status,
            http_status=response.status_code,
            discovered_count=len(candidates),
            new_count=new_count,
            changed_count=changed_count,
            known_count=known_count,
            message=message,
        )
    except Exception as exc:
        # PostgreSQL leaves the transaction unusable after an error such as a
        # check-constraint violation. Roll back before writing the failed run
        # result and notification so the admin sees a controlled failure instead
        # of a secondary InFailedSqlTransaction error.
        conn.rollback()
        message = "Source check failed. No application data was changed."
        finished = repo.finish_check_run(
            conn,
            check_run_id=check_run_id,
            status="failed",
            http_status=None,
            discovered_count=0,
            new_count=0,
            changed_count=0,
            known_count=0,
            message=message,
            error_message=str(exc),
        )
        repo.create_notification(
            conn,
            notification_type="source_check_failed",
            severity="error",
            title="MYH source check failed",
            message=str(exc),
            actor_user_id=actor_user_id,
        )

    return {
        "check_run_id": finished["id"],
        "status": finished["status"],
        "source_url": finished["source_url"],
        "started_at": finished["started_at"],
        "finished_at": finished.get("finished_at"),
        "http_status": finished.get("http_status"),
        "discovered_count": finished.get("discovered_count") or 0,
        "new_count": finished.get("new_count") or 0,
        "changed_count": finished.get("changed_count") or 0,
        "known_count": finished.get("known_count") or 0,
        "message": finished.get("message") or message,
        "error_message": finished.get("error_message"),
        "files": files,
    }


def relative_runtime_path(path: Path) -> str:
    """Return a stable relative path where possible."""
    resolved = path.resolve()
    for base in (BACKEND_ROOT.parent, BACKEND_ROOT, Path.cwd()):
        try:
            return resolved.relative_to(base.resolve()).as_posix()
        except ValueError:
            continue
    return resolved.as_posix()


def _safe_file_name(file_name: str) -> str:
    """Return a filesystem-safe file name while keeping the extension."""
    cleaned = re.sub(r"[^A-Za-z0-9._-]+", "_", file_name).strip("._")
    return cleaned or "myh_source_file"


def sha256_bytes(payload: bytes) -> str:
    """Return SHA256 hex digest for downloaded bytes."""
    return hashlib.sha256(payload).hexdigest()


def download_source_file(
    conn: psycopg.Connection[dict[str, Any]],
    *,
    source_file_id: str,
    config: SourceMonitorConfig | None = None,
    downloader: Callable[[str, int], bytes] = download_url_bytes,
) -> dict[str, Any]:
    """Download a stored source file and persist SHA256/runtime path metadata."""
    config = config or build_config_from_env()
    source_file = repo.get_source_file(conn, source_file_id)
    if source_file is None:
        raise FileNotFoundError(f"Source file {source_file_id} was not found.")

    config.download_dir.mkdir(parents=True, exist_ok=True)
    try:
        payload = downloader(source_file["file_url"], config.timeout_seconds)
        digest = sha256_bytes(payload)
        file_name = _safe_file_name(source_file["file_name"])
        target = config.download_dir / f"{source_file_id}_{file_name}"
        target.write_bytes(payload)
        changed = bool(source_file.get("last_sha256") and source_file.get("last_sha256") != digest)
        stored = repo.update_download_metadata(
            conn,
            source_file_id=source_file_id,
            sha256=digest,
            downloaded_path=relative_runtime_path(target),
            size_bytes=len(payload),
            changed=changed,
        )
        if changed:
            repo.create_notification(
                conn,
                notification_type="source_file_content_changed",
                severity="warning",
                title="Downloaded MYH source content changed",
                message=f"{source_file['file_name']} has a new SHA256 hash and should be reviewed before import.",
                source_file_id=source_file_id,
            )
        return {
            "source_file_id": stored["id"],
            "file_name": stored["file_name"],
            "downloaded_path": stored["downloaded_path"],
            "sha256": digest,
            "size_bytes": len(payload),
            "changed": changed,
            "message": "Source file downloaded and hashed.",
        }
    except Exception as exc:
        repo.update_source_file_error(conn, source_file_id=source_file_id, error=str(exc))
        raise


def _normalize_header(value: Any) -> str:
    text = " ".join(str(value or "").strip().lower().split())
    return text.replace("_", " ")


def _build_header_map(headers: Iterable[Any]) -> dict[str, int]:
    normalized_headers = [_normalize_header(header) for header in headers]
    mapping: dict[str, int] = {}
    for canonical, aliases in HEADER_ALIASES.items():
        normalized_aliases = {_normalize_header(alias) for alias in aliases}
        for idx, header in enumerate(normalized_headers):
            if header in normalized_aliases or any(alias == header for alias in normalized_aliases):
                mapping[canonical] = idx
                break
    return mapping


def _sheet_rows_from_csv(path: Path) -> tuple[str, list[list[Any]]]:
    with path.open("r", encoding="utf-8-sig", newline="") as file:
        return path.stem, list(csv.reader(file))


def _sheet_rows_from_excel(path: Path) -> tuple[str, list[list[Any]]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("openpyxl is required to import MYH Excel source files.") from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    preferred_names = [name for name in workbook.sheetnames if "tabell 3" in name.lower()]
    candidate_names = preferred_names or list(workbook.sheetnames)
    best: tuple[str, list[list[Any]], int] | None = None
    for sheet_name in candidate_names:
        sheet = workbook[sheet_name]
        rows = [list(row) for row in sheet.iter_rows(values_only=True)]
        score = 0
        for row in rows[:80]:
            mapping = _build_header_map(row)
            score = max(score, len(REQUIRED_RAW_FIELDS & set(mapping)))
        if best is None or score > best[2]:
            best = (sheet_name, rows, score)
    if best is None:
        raise ValueError("The workbook does not contain readable sheets.")
    return best[0], best[1]


def _read_source_table(path: Path) -> tuple[str, list[list[Any]]]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return _sheet_rows_from_csv(path)
    if suffix in {".xlsx", ".xlsm", ".xls"}:
        return _sheet_rows_from_excel(path)
    raise ValueError(f"Unsupported source file type: {suffix}")


def _find_header(rows: list[list[Any]]) -> tuple[int, dict[str, int]]:
    best_index = -1
    best_map: dict[str, int] = {}
    best_score = 0
    for idx, row in enumerate(rows[:100]):
        mapping = _build_header_map(row)
        score = len(REQUIRED_RAW_FIELDS & set(mapping))
        if score > best_score:
            best_index = idx
            best_map = mapping
            best_score = score
    if best_score < 5:
        raise ValueError("Could not find a Tabell 3-like header row in the official source file.")
    missing = sorted(REQUIRED_RAW_FIELDS - set(best_map))
    if missing:
        raise ValueError(f"Official source file is missing required Tabell 3 columns: {missing}")
    return best_index, best_map


def _cell(row: list[Any], mapping: dict[str, int], key: str) -> str | None:
    index = mapping.get(key)
    if index is None or index >= len(row):
        return None
    value = row[index]
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _required_cell(row: list[Any], mapping: dict[str, int], key: str) -> str:
    value = clean_text(_cell(row, mapping, key))
    if value is None:
        raise ValueError(f"Required source field {key!r} is blank.")
    return value


def normalize_decision(value: str) -> str:
    """Normalize Swedish MYH decision labels to the curated decision codes."""
    normalized = value.strip().lower()
    if "bevilj" in normalized and "ej" not in normalized:
        return "approved"
    if "ej bevil" in normalized or "avslag" in normalized or "avslagen" in normalized:
        return "rejected"
    if "åter" in normalized or "ater" in normalized or "withdraw" in normalized:
        return "withdrawn"
    raise ValueError(f"Unexpected decision value: {value!r}")


def _normalized_principal_type(value: str) -> str:
    text = value.strip().lower()
    if "kommun" in text:
        return "municipal"
    if "region" in text or "landsting" in text:
        return "regional"
    if "stat" in text:
        return "state"
    if "priv" in text or "enskild" in text:
        return "private"
    return "other"


def _to_int_text(value: str | None, default: int) -> str:
    cleaned = clean_text(value)
    if cleaned is None:
        return str(default)
    cleaned = cleaned.replace("%", "").replace(",", ".")
    try:
        return str(int(float(cleaned)))
    except ValueError:
        return str(default)


def _bool_text(value: bool) -> str:
    return "true" if value else "false"


def _build_curated_row(
    raw_row: list[Any],
    *,
    mapping: dict[str, int],
    source_year: int,
    source_file: str,
    source_sheet: str,
    source_row: int,
) -> dict[str, str]:
    beslut = _required_cell(raw_row, mapping, "beslut")
    decision_code = normalize_decision(beslut)
    studieform = _required_cell(raw_row, mapping, "studieform")
    kommun = _required_cell(raw_row, mapping, "kommun")
    multiple = "," in kommun or ";" in kommun or clean_text(_cell(raw_row, mapping, "flera_kommuner")) in {"Ja", "ja", "True", "true"}
    return {
        "source_year": str(source_year),
        "source_file": source_file,
        "source_sheet": source_sheet,
        "source_row": str(source_row),
        "diarienummer": _required_cell(raw_row, mapping, "diarienummer"),
        "utbildningsnamn": _required_cell(raw_row, mapping, "utbildningsnamn"),
        "utbildningsomrade": _required_cell(raw_row, mapping, "utbildningsomrade"),
        "beslut": beslut,
        "beslut_normalized": decision_code,
        "is_approved": _bool_text(decision_code == "approved"),
        "lan": _required_cell(raw_row, mapping, "lan"),
        "kommun": kommun,
        "flera_kommuner": "Ja" if multiple else "Nej",
        "has_multiple_municipalities": _bool_text(multiple),
        "antal_kommuner": _to_int_text(_cell(raw_row, mapping, "antal_kommuner"), 2 if multiple else 1),
        "yh_poang": _to_int_text(_cell(raw_row, mapping, "yh_poang"), 1),
        "studieform": studieform,
        "is_distance_based": _bool_text("distans" in studieform.lower()),
        "studietakt_procent": _to_int_text(_cell(raw_row, mapping, "studietakt_procent"), 100),
        "examenstyp": clean_text(_cell(raw_row, mapping, "examenstyp")) or "",
        "utbildningsanordnare": _required_cell(raw_row, mapping, "utbildningsanordnare"),
        "huvudmannatyp": _required_cell(raw_row, mapping, "huvudmannatyp"),
        "huvudmannatyp_normalized": _normalized_principal_type(_required_cell(raw_row, mapping, "huvudmannatyp")),
        "sokta_utbildningsomgangar": _to_int_text(_cell(raw_row, mapping, "sokta_utbildningsomgangar"), 1),
        "beviljade_utbildningsomgangar": _to_int_text(_cell(raw_row, mapping, "beviljade_utbildningsomgangar"), 1 if decision_code == "approved" else 0),
        "sun5_inriktning": clean_text(_cell(raw_row, mapping, "sun5_inriktning")) or "",
        "sun5_inriktning_namn": clean_text(_cell(raw_row, mapping, "sun5_inriktning_namn")) or "",
        "seqf_niva": clean_text(_cell(raw_row, mapping, "seqf_niva")) or "",
        "smalt_yrkesomrade": clean_text(_cell(raw_row, mapping, "smalt_yrkesomrade")) or "",
        "sokta_platser_per_utbildningsomgang": clean_text(_cell(raw_row, mapping, "sokta_platser_per_utbildningsomgang")) or "",
        "sokta_platser_totalt": clean_text(_cell(raw_row, mapping, "sokta_platser_totalt")) or "",
        "beviljade_platser_totalt": clean_text(_cell(raw_row, mapping, "beviljade_platser_totalt")) or "",
    }


def transform_official_source_file(path: Path, *, source_year: int | None = None, source_file_name: str | None = None) -> list[dict[str, str]]:
    """Transform one official MYH source file to curated application rows."""
    sheet_name, rows = _read_source_table(path)
    header_index, mapping = _find_header(rows)
    resolved_year = source_year or extract_source_year(source_file_name or path.name, path.name)
    if resolved_year is None:
        raise ValueError("Could not determine source year from file metadata. Store source_year before importing.")

    curated_rows: list[dict[str, str]] = []
    for physical_index, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
        if not any(clean_text(str(value)) for value in row if value is not None):
            continue
        try:
            curated_rows.append(
                _build_curated_row(
                    row,
                    mapping=mapping,
                    source_year=resolved_year,
                    source_file=source_file_name or path.name,
                    source_sheet=sheet_name,
                    source_row=physical_index,
                )
            )
        except ValueError as exc:
            raise ValueError(f"Invalid source row {physical_index}: {exc}") from exc
    validate_curated_import_rows(curated_rows)
    return curated_rows


def validate_curated_import_rows(rows: list[dict[str, str]]) -> None:
    """Validate transformed official source rows before database import."""
    if not rows:
        raise ValueError("The official source file produced zero importable Tabell 3 rows.")
    diarienummer_values = [row["diarienummer"] for row in rows]
    if len(diarienummer_values) != len(set(diarienummer_values)):
        raise ValueError("diarienummer is not unique in the transformed official source rows.")
    years = {to_int(row.get("source_year"), "source_year") for row in rows}
    if any(year < 2020 or year > 2100 for year in years):
        raise ValueError(f"Source years must be between 2020 and 2100. Found {sorted(years)}.")
    decisions = {row["beslut_normalized"] for row in rows}
    unexpected = decisions - {"approved", "rejected", "withdrawn"}
    if unexpected:
        raise ValueError(f"Unexpected normalized decision values: {sorted(unexpected)}")


def _fetch_id_maps(conn: psycopg.Connection[dict[str, Any]]) -> dict[str, dict[Any, int]]:
    """Fetch lookup IDs using dict-row-safe access."""

    def get(row: Any, key: str, index: int) -> Any:
        if isinstance(row, dict):
            return row[key]
        return row[index]

    maps: dict[str, dict[Any, int]] = {}
    with conn.cursor() as cursor:
        cursor.execute("SELECT provider_id, utbildningsanordnare FROM providers;")
        maps["providers"] = {get(row, "utbildningsanordnare", 1): get(row, "provider_id", 0) for row in cursor.fetchall()}
        cursor.execute("SELECT education_area_id, utbildningsomrade FROM education_areas;")
        maps["education_areas"] = {get(row, "utbildningsomrade", 1): get(row, "education_area_id", 0) for row in cursor.fetchall()}
        cursor.execute("SELECT location_id, lan, kommun FROM locations;")
        maps["locations"] = {(get(row, "lan", 1), get(row, "kommun", 2)): get(row, "location_id", 0) for row in cursor.fetchall()}
        cursor.execute("SELECT principal_type_id, huvudmannatyp FROM principal_types;")
        maps["principal_types"] = {get(row, "huvudmannatyp", 1): get(row, "principal_type_id", 0) for row in cursor.fetchall()}
        cursor.execute("SELECT study_form_id, studieform FROM study_forms;")
        maps["study_forms"] = {get(row, "studieform", 1): get(row, "study_form_id", 0) for row in cursor.fetchall()}
    return maps


def _load_application_rows(conn: psycopg.Connection[dict[str, Any]], rows: list[dict[str, str]]) -> None:
    """Insert transformed application rows using existing curated-data mapping rules."""
    id_maps = _fetch_id_maps(conn)
    application_rows = [build_application_row(row, id_maps) for row in rows]
    with conn.cursor() as cursor:
        cursor.executemany(APPLICATION_INSERT_SQL, application_rows)


def _write_processed_csv(rows: list[dict[str, str]], processed_dir: Path, refresh_run_id: str) -> Path:
    processed_dir.mkdir(parents=True, exist_ok=True)
    target = processed_dir / f"refresh_{refresh_run_id}.csv"
    fieldnames = list(rows[0].keys())
    with target.open("w", encoding="utf-8", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)
    return target


def import_transformed_rows(
    conn: psycopg.Connection[dict[str, Any]],
    *,
    rows: list[dict[str, str]],
    source_file_id: str | None,
    refresh_run_id: str,
    mode: str,
    source_sha256: str | None,
    downloaded_path: str | None,
    processed_dir: Path,
) -> dict[str, Any]:
    """Atomically replace only affected official source years."""
    validate_curated_import_rows(rows)
    affected_years = sorted({to_int(row.get("source_year"), "source_year") for row in rows})
    processed_path = _write_processed_csv(rows, processed_dir, refresh_run_id)
    validation_summary = f"Validated {len(rows)} rows for source year(s): {', '.join(map(str, affected_years))}."

    with conn.transaction():
        insert_lookup_rows(conn, rows)
        with conn.cursor() as cursor:
            cursor.execute("DELETE FROM applications WHERE source_year = ANY(%(years)s);", {"years": affected_years})
        _load_application_rows(conn, rows)

    return repo.finish_refresh_run(
        conn,
        refresh_run_id=refresh_run_id,
        status="success",
        rows_imported=len(rows),
        affected_years=",".join(map(str, affected_years)),
        source_sha256=source_sha256,
        downloaded_path=downloaded_path,
        processed_path=relative_runtime_path(processed_path),
        validation_summary=validation_summary,
    )


def import_source_file(
    conn: psycopg.Connection[dict[str, Any]],
    *,
    source_file_id: str,
    actor_user_id: str | None,
    config: SourceMonitorConfig | None = None,
) -> dict[str, Any]:
    """Import a downloaded official source file into the applications table."""
    config = config or build_config_from_env()
    source_file = repo.get_source_file(conn, source_file_id)
    if source_file is None:
        raise FileNotFoundError(f"Source file {source_file_id} was not found.")
    if not source_file.get("downloaded_path"):
        raise ValueError("Download the official source file before importing it.")

    refresh_run = repo.create_refresh_run(conn, mode="official_source_file", source_file_id=source_file_id, triggered_by_user_id=actor_user_id)
    try:
        downloaded_path = Path(str(source_file["downloaded_path"]))
        if not downloaded_path.is_absolute():
            downloaded_path = BACKEND_ROOT.parent / downloaded_path
        rows = transform_official_source_file(downloaded_path, source_year=source_file.get("source_year"), source_file_name=source_file["file_name"])
        finished = import_transformed_rows(
            conn,
            rows=rows,
            source_file_id=source_file_id,
            refresh_run_id=str(refresh_run["id"]),
            mode="official_source_file",
            source_sha256=source_file.get("last_sha256"),
            downloaded_path=source_file.get("downloaded_path"),
            processed_dir=config.processed_dir,
        )
        repo.create_notification(
            conn,
            notification_type="refresh_succeeded",
            severity="success",
            title="Official MYH refresh succeeded",
            message=f"Imported {finished.get('rows_imported')} row(s) from {source_file['file_name']}.",
            source_file_id=source_file_id,
            refresh_run_id=finished["id"],
            actor_user_id=actor_user_id,
        )
        return finished
    except Exception as exc:
        finished = repo.finish_refresh_run(
            conn,
            refresh_run_id=refresh_run["id"],
            status="failed",
            error_message=str(exc),
            downloaded_path=source_file.get("downloaded_path"),
            source_sha256=source_file.get("last_sha256"),
        )
        repo.create_notification(
            conn,
            notification_type="refresh_failed",
            severity="error",
            title="Official MYH refresh failed",
            message=str(exc),
            source_file_id=source_file_id,
            refresh_run_id=finished["id"],
            actor_user_id=actor_user_id,
        )
        raise


def refresh_from_curated_csv_compatibility(
    conn: psycopg.Connection[dict[str, Any]],
    *,
    actor_user_id: str | None,
    config: SourceMonitorConfig | None = None,
    csv_path: Path | None = None,
) -> dict[str, Any]:
    """Compatibility refresh for POST /refresh through the run/notification service.

    This keeps the legacy endpoint useful for local demos while preserving the
    safer 3.20 behavior: admin auth, refresh-run history, validation, atomic
    affected-year replacement, and no destructive database reset.
    """
    config = config or build_config_from_env()
    refresh_run = repo.create_refresh_run(conn, mode="curated_csv_compatibility", source_file_id=None, triggered_by_user_id=actor_user_id)
    try:
        resolved_csv = resolve_csv_path(csv_path)
        rows = read_rows(resolved_csv)
        finished = import_transformed_rows(
            conn,
            rows=rows,
            source_file_id=None,
            refresh_run_id=str(refresh_run["id"]),
            mode="curated_csv_compatibility",
            source_sha256=sha256_bytes(resolved_csv.read_bytes()),
            downloaded_path=relative_runtime_path(resolved_csv),
            processed_dir=config.processed_dir,
        )
        repo.create_notification(
            conn,
            notification_type="refresh_succeeded",
            severity="success",
            title="Compatibility refresh succeeded",
            message=f"Reloaded {finished.get('rows_imported')} official curated row(s) from the local CSV.",
            refresh_run_id=finished["id"],
            actor_user_id=actor_user_id,
        )
        return finished
    except Exception as exc:
        finished = repo.finish_refresh_run(conn, refresh_run_id=refresh_run["id"], status="failed", error_message=str(exc))
        repo.create_notification(
            conn,
            notification_type="refresh_failed",
            severity="error",
            title="Compatibility refresh failed",
            message=str(exc),
            refresh_run_id=finished["id"],
            actor_user_id=actor_user_id,
        )
        raise


def build_status(conn: psycopg.Connection[dict[str, Any]], *, config: SourceMonitorConfig | None = None) -> dict[str, Any]:
    """Build the admin source-monitor status payload."""
    config = config or build_config_from_env()
    return {
        "monitor_enabled": config.monitor_enabled,
        "run_on_startup": config.run_on_startup,
        "auto_import_enabled": config.auto_import_enabled,
        "interval_minutes": config.interval_minutes,
        "source_url": config.source_url,
        "last_check": repo.latest_check_run(conn),
        "known_source_files": repo.count_source_files(conn),
        "unread_notifications": repo.count_unread_notifications(conn),
        "latest_refresh_run": repo.latest_refresh_run(conn),
    }


def run_scheduled_check_once(config: SourceMonitorConfig | None = None) -> None:
    """Run a scheduler-owned source check with its own database connection."""
    config = config or build_config_from_env()
    with open_connection(get_database_url()) as conn:
        result = run_source_check(conn, config=config)
        if config.auto_import_enabled:
            for source_file in result.get("files", []):
                if source_file.get("detection_status") not in {"new", "changed"}:
                    continue
                download_source_file(conn, source_file_id=str(source_file["id"]), config=config)
                import_source_file(conn, source_file_id=str(source_file["id"]), actor_user_id=None, config=config)


def copy_fixture_to_download_dir(source_path: Path, target_dir: Path) -> Path:
    """Small helper used by focused tests to avoid live network downloads."""
    target_dir.mkdir(parents=True, exist_ok=True)
    target = target_dir / source_path.name
    shutil.copyfile(source_path, target)
    return target
